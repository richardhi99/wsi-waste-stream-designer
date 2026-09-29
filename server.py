#!/usr/bin/env python3
"""Water Services Inc. customer intake server.

Serves the two-page site and writes each request to an Excel workbook
that Microsoft Excel and Access can open.

  python3 server.py

Customer page:  http://127.0.0.1:8765/customer.html
Office queue:   http://127.0.0.1:8765/office.html
Workbook:       data/WSI_Customer_Submissions.xlsx

Optional SMTP (sends the customer thank-you from Water Services Inc):
  export WSI_SMTP_HOST=smtp.office365.com
  export WSI_SMTP_PORT=587
  export WSI_SMTP_USER=info@water-services.us
  export WSI_SMTP_PASS=...
  export WSI_FROM=info@water-services.us
"""
from __future__ import annotations

import json
import os
import smtplib
import uuid
from datetime import datetime
from email.message import EmailMessage
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
XLSX = DATA / "WSI_Customer_Submissions.xlsx"
JSONL = DATA / "submissions.jsonl"
HOST = os.environ.get("WSI_HOST", "0.0.0.0")
PORT = int(os.environ.get("PORT") or os.environ.get("WSI_PORT", "8765"))

HEADERS = [
    "ID", "Submitted", "Project", "Client", "Contact", "Email", "Phone",
    "Stream", "Goal", "ADF_MGD", "Peak_factor", "Tank_type",
    "Min_depth_ft", "Max_depth_ft", "Aerator", "Train",
    "COG_units", "COG_O2_lb_d", "ISO_boxes", "ISO_list",
    "Biology", "AOR_lb_d", "Notes", "JSON",
]


def _xlsx():
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    DATA.mkdir(exist_ok=True)
    if XLSX.exists():
        return load_workbook(XLSX)
    wb = Workbook()
    ws = wb.active
    ws.title = "Submissions"
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1B365D")
    for col, name in enumerate(HEADERS, 1):
        cell = ws.cell(1, col, name)
        cell.font = head
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True)
    from openpyxl.utils import get_column_letter
    widths = [16, 20, 22, 22, 18, 28, 16, 22, 24, 12, 12, 22, 12, 12, 22, 40, 24, 12, 10, 28, 22, 12, 36, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    info = wb.create_sheet("Readme")
    info["A1"] = "Water Services Inc. customer request database"
    info["A2"] = "Excel / Microsoft Access compatible. Each row is one web-form submission."
    info["A3"] = "Office page: office.html   Customer page: customer.html"
    info["A4"] = "372 South 900 West, Provo, Utah 84601  ·  info@water-services.us"
    return wb


def append_row(payload: dict) -> str:
    rec_id = payload.get("id") or ("WSI-" + datetime.now().strftime("%Y%m%d-") + uuid.uuid4().hex[:6].upper())
    payload["id"] = rec_id
    i = payload.get("inputs") or {}
    sm = payload.get("summary") or {}
    row = [
        rec_id,
        payload.get("submitted_at") or datetime.now().isoformat(timespec="seconds"),
        i.get("project"), i.get("client"), i.get("contact"), i.get("email"), i.get("phone"),
        i.get("stream"), i.get("goal"), i.get("adf"), i.get("pf"),
        i.get("tankType"), i.get("minDepth"), i.get("maxDepth"), i.get("aerator"),
        sm.get("train"), sm.get("cog"), sm.get("cog_o2_lb_d"),
        sm.get("containers"), sm.get("container_list"),
        sm.get("family"), sm.get("aor_lb"), i.get("note"),
        json.dumps(payload, default=str),
    ]
    DATA.mkdir(exist_ok=True)
    with JSONL.open("a", encoding="utf-8") as f:
        f.write(json.dumps(payload, default=str) + "\n")
    try:
        wb = _xlsx()
        ws = wb["Submissions"]
        ws.append(row)
        wb.save(XLSX)
    except Exception as exc:
        print("Excel write failed:", exc)
    return rec_id


def list_submissions() -> list:
    rows = []
    if JSONL.exists():
        for line in JSONL.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line:
                continue
            try:
                rows.append(json.loads(line))
            except json.JSONDecodeError:
                continue
    rows.reverse()
    return rows


def send_thanks(to_addr: str, name: str, rec_id: str) -> bool:
    host = os.environ.get("WSI_SMTP_HOST")
    user = os.environ.get("WSI_SMTP_USER")
    pw = os.environ.get("WSI_SMTP_PASS")
    if not (host and user and pw and to_addr):
        return False
    port = int(os.environ.get("WSI_SMTP_PORT", "587"))
    sender = os.environ.get("WSI_FROM", user)
    msg = EmailMessage()
    msg["Subject"] = "We received your request — Water Services Inc"
    msg["From"] = sender
    msg["To"] = to_addr
    who = name or "there"
    msg.set_content(
        f"Hello {who},\n\n"
        "Thank you. We received the waste-stream information you submitted.\n\n"
        "Water Services Inc will get back to you.\n\n"
        f"Reference: {rec_id}\n\n"
        "Water Services, Inc.\n"
        "372 South 900 West, Provo, Utah 84601\n"
        "+1-801-705-4567\n"
        "info@water-services.us\n"
        "www.water-services.us\n"
    )
    try:
        with smtplib.SMTP(host, port, timeout=20) as smtp:
            smtp.starttls()
            smtp.login(user, pw)
            smtp.send_message(msg)
        return True
    except Exception as exc:
        print("SMTP failed:", exc)
        return False


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT), **kwargs)

    def end_headers(self):
        self.send_header("Cache-Control", "no-store")
        super().end_headers()

    def do_GET(self):
        path = urlparse(self.path).path
        if path in ("/", "/index.html"):
            self.path = "/customer.html"
            return SimpleHTTPRequestHandler.do_GET(self)
        if path == "/api/submissions":
            body = json.dumps(list_submissions(), default=str).encode()
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return
        if path == "/api/database.xlsx":
            if not XLSX.exists():
                _xlsx().save(XLSX)
            data = XLSX.read_bytes()
            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition", "attachment; filename=WSI_Customer_Submissions.xlsx")
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)
            return
        return SimpleHTTPRequestHandler.do_GET(self)

    def do_POST(self):
        path = urlparse(self.path).path
        if path != "/api/submit":
            self.send_error(404)
            return
        length = int(self.headers.get("Content-Length", "0"))
        raw = self.rfile.read(length)
        try:
            payload = json.loads(raw.decode("utf-8"))
        except json.JSONDecodeError:
            self.send_error(400, "invalid json")
            return
        rec_id = append_row(payload)
        emailed = send_thanks(
            (payload.get("inputs") or {}).get("email") or "",
            (payload.get("inputs") or {}).get("contact") or "",
            rec_id,
        )
        body = json.dumps({"ok": True, "id": rec_id, "emailed": emailed}).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, fmt, *args):
        print("[%s] %s" % (self.log_date_time_string(), fmt % args))


if __name__ == "__main__":
    DATA.mkdir(exist_ok=True)
    if not XLSX.exists():
        _xlsx().save(XLSX)
    httpd = ThreadingHTTPServer((HOST, PORT), Handler)
    print(f"Customer page  http://{HOST}:{PORT}/customer.html")
    print(f"Office queue   http://{HOST}:{PORT}/office.html")
    print(f"Database       {XLSX}")
    httpd.serve_forever()
